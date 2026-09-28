"""Small OOXML fixtures for the office craft gate."""
from __future__ import annotations

import sys
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'plugins/litcodex/skills/lit-pptx/scripts'))
from craft_extras import analyze_craft  # noqa: E402


def deck():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(7), Inches(.5))
    title.text = 'Operations review'
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    return prs, slide


def shape(slide, x, y, w, h, color=(44, 86, 150), kind=MSO_SHAPE.RECTANGLE):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = RGBColor(*color)
    return item


def text(slide, value, x=.8, y=1.5, w=4, h=.5, size=16):
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    item.text = value
    for paragraph in item.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
    return item


def severity(prs, rule):
    return [finding['severity'] for finding in analyze_craft(prs)['findings'] if finding['rule'] == rule]


class PptxCraftChecks(unittest.TestCase):
    def test_accent_families(self):
        poor, slide = deck()
        for x, rgb in [(1, (205, 51, 51)), (2, (48, 158, 76)), (3, (54, 76, 190))]:
            shape(slide, x, 1.7, .5, .5, rgb)
        clean, slide = deck()
        shape(slide, 1, 1.7, .5, .5)
        self.assertIn('HIGH', severity(poor, 'OF-101'))
        self.assertNotIn('HIGH', severity(clean, 'OF-101'))

    def test_body_measure(self):
        poor, slide = deck()
        text(slide, 'A' * 100 + '\n' + 'B' * 100, w=11, h=1, size=11)
        clean, slide = deck()
        text(slide, 'A useful short line\nAnother short line', w=5, h=1)
        self.assertIn('HIGH', severity(poor, 'OF-102'))
        self.assertNotIn('HIGH', severity(clean, 'OF-102'))

    def test_numeric_table_alignment(self):
        def fixture(alignment):
            prs, slide = deck()
            table = slide.shapes.add_table(3, 2, Inches(1), Inches(1.5), Inches(5), Inches(2)).table
            table.cell(0, 0).text, table.cell(0, 1).text = 'Period', 'Amount'
            for row, value in enumerate(('1,200', '2,400'), 1):
                table.cell(row, 0).text = f'Q{row}'
                table.cell(row, 1).text = value
                table.cell(row, 1).text_frame.paragraphs[0].alignment = alignment
            return prs
        self.assertIn('HIGH', severity(fixture(PP_ALIGN.LEFT), 'OF-103'))
        self.assertNotIn('HIGH', severity(fixture(PP_ALIGN.RIGHT), 'OF-103'))

    def test_concentric_radius(self):
        def fixture(adjustment):
            prs, slide = deck()
            shape(slide, 1, 1.5, 4, 2, kind=MSO_SHAPE.ROUNDED_RECTANGLE)
            inner = shape(slide, 1.2, 1.7, 3.6, 1.6, kind=MSO_SHAPE.ROUNDED_RECTANGLE)
            inner.adjustments[0] = adjustment
            return prs
        self.assertIn('HIGH', severity(fixture(.16667), 'OF-104'))
        self.assertNotIn('HIGH', severity(fixture(.08333), 'OF-104'))

    def test_group_gap_ratio(self):
        def fixture(second_x):
            prs, slide = deck()
            shape(slide, .7, 1.4, 2, 1.3)
            shape(slide, second_x, 1.4, 2, 1.3)
            text(slide, 'First', .85, 1.6, .6, .4)
            text(slide, 'Second', 1.65, 1.6, .7, .4)
            text(slide, 'Third', second_x + .15, 1.6, .6, .4)
            text(slide, 'Fourth', second_x + .95, 1.6, .7, .4)
            return prs
        self.assertIn('MEDIUM', severity(fixture(2.8), 'OF-105'))
        self.assertNotIn('MEDIUM', severity(fixture(4.2), 'OF-105'))

    def test_gradient_text(self):
        poor, slide = deck()
        run = text(slide, 'Plain heading').text_frame.paragraphs[0].runs[0]
        run._r.get_or_add_rPr().append(OxmlElement('a:gradFill'))
        clean, slide = deck()
        text(slide, 'Plain heading')
        self.assertIn('HIGH', severity(poor, 'OF-106'))
        self.assertNotIn('HIGH', severity(clean, 'OF-106'))

    def test_glow_effect(self):
        poor, slide = deck()
        item = shape(slide, 1, 1.5, 2, 1)
        effect = OxmlElement('a:effectLst')
        effect.append(OxmlElement('a:glow'))
        item._element.spPr.append(effect)
        clean, slide = deck()
        shape(slide, 1, 1.5, 2, 1)
        self.assertIn('HIGH', severity(poor, 'OF-107'))
        self.assertNotIn('HIGH', severity(clean, 'OF-107'))

    def test_emoji_bullet(self):
        poor, slide = deck()
        paragraph = text(slide, 'Task item').text_frame.paragraphs[0]
        bullet = OxmlElement('a:buChar')
        bullet.set('char', '🚀')
        paragraph._p.get_or_add_pPr().append(bullet)
        clean, slide = deck()
        text(slide, 'Task item')
        self.assertIn('HIGH', severity(poor, 'OF-108'))
        self.assertNotIn('HIGH', severity(clean, 'OF-108'))

    def test_empty_trailing_band(self):
        poor, slide = deck()
        text(slide, 'Short block', y=1.5)
        clean, slide = deck()
        for y in (1.5, 3.1, 4.7, 6.0):
            text(slide, 'Content block', y=y)
        self.assertIn('HIGH', severity(poor, 'OF-109'))
        self.assertNotIn('HIGH', severity(clean, 'OF-109'))


if __name__ == '__main__':
    unittest.main()
