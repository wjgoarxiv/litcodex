"""Small real OOXML regressions for the bundled office engines.

Run with the office runner's pinned Python environment.
"""

from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from bs4 import BeautifulSoup
from docx import Document
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1] / "plugins/litcodex/skills"


def load(name, path):
    spec = spec_from_file_location(name, path)
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


pptx_qa = load("pptx_qa", ROOT / "lit-pptx/scripts/qa_deck.py")
docx_converter = load("docx_converter", ROOT / "lit-docx/scripts/convert_md_to_docx.py")


class OfficeQualityTests(unittest.TestCase):
    def test_sparse_content_slide_has_actionable_failure(self):
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11), Inches(0.7))
        title.text = "Quarterly results"
        body = slide.shapes.add_textbox(Inches(0.8), Inches(2), Inches(4), Inches(0.4))
        body.text = "One short point"
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "sparse.pptx"
            deck.save(path)
            result = pptx_qa.check_content_quality(path)
        self.assertFalse(result["pass"])
        self.assertIn("fill", str(result).lower())
        self.assertIn("fix", str(result).lower())

    def test_shallow_kpi_cards_do_not_count_as_a_filled_content_area(self):
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11), Inches(0.7)).text = "Headline"
        for x in (0.8, 6.8):
            card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.05), Inches(5.7), Inches(1.75))
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(230, 239, 255)
            slide.shapes.add_textbox(Inches(x + 0.2), Inches(2.3), Inches(5.0), Inches(1.0)).text = "205"
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "shallow-kpi.pptx"
            deck.save(path)
            result = pptx_qa.check_content_quality(path)
        self.assertFalse(result["pass"])
        self.assertIn("underfilled", str(result))

    def test_table_only_deck_has_actionable_failure(self):
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        for _ in range(2):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11), Inches(0.7)).text = "Results"
            table = slide.shapes.add_table(4, 2, Inches(0.8), Inches(2), Inches(11), Inches(3)).table
            table.cell(0, 0).text, table.cell(0, 1).text = "Period", "Value"
            for row in range(1, 4):
                table.cell(row, 0).text, table.cell(row, 1).text = f"Q{row}", str(row * 12)
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "tables.pptx"
            deck.save(path)
            result = pptx_qa.check_content_quality(path)
        self.assertFalse(result["pass"])
        self.assertIn("table-only", str(result).lower())

    def test_unfilled_cover_box_is_rejected(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text = "Title"
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(2), Inches(4), Inches(0.4), Inches(0.4))
        box.fill.background()
        box.line.color.rgb = RGBColor(0, 0, 0)
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "cover.pptx"
            deck.save(path)
            result = pptx_qa.check_content_quality(path)
        self.assertFalse(result["pass"])
        self.assertIn("empty-shape", str(result))

    def test_chart_overlaid_by_body_text_is_rejected(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(8), Inches(0.7)).text = "Results"
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        data = CategoryChartData()
        data.categories = ["Q1", "Q2"]
        data.add_series("Value", [10, 20])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(10), Inches(3), data)
        slide.shapes.add_textbox(Inches(2), Inches(3), Inches(4), Inches(0.5)).text = "Text over plot"
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "occluded-chart.pptx"
            deck.save(path)
            result = pptx_qa.check_content_quality(path)
        self.assertFalse(result["pass"])
        self.assertIn("chart-overlap", str(result))

    def test_docx_table_has_fixed_width_and_unsplit_rows(self):
        doc = Document()
        html = BeautifulSoup("<table><tr><th>Item</th><th>Description</th></tr>"
                             "<tr><td>Phase</td><td>A long descriptive cell that should receive most of the width.</td></tr></table>",
                             "html.parser")
        docx_converter.process_table(doc, html.table, Path("sample.md"))
        table = doc.tables[0]
        self.assertFalse(table.autofit)
        self.assertGreater(table.columns[1].width, table.columns[0].width)
        self.assertTrue(all(row._tr.trPr.find(docx_converter.qn("w:cantSplit")) is not None for row in table.rows))
        self.assertIsNotNone(table.rows[0]._tr.trPr.find(docx_converter.qn("w:tblHeader")))

    def test_publisher_promotes_first_heading_to_title(self):
        with TemporaryDirectory() as tmp:
            source, output = Path(tmp) / "source.md", Path(tmp) / "result.docx"
            source.write_text("# Sample plan\n\n## Summary\n\nAn illustrative example.\n")
            docx_converter.convert_md_to_docx(source, output, publisher_name="korean-generic")
            paragraphs = [p for p in Document(output).paragraphs if p.text.strip()]
        self.assertEqual(paragraphs[0].text, "Sample plan")
        self.assertEqual(paragraphs[0].alignment, docx_converter.WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(sum(p.text == "Sample plan" for p in paragraphs), 1)


if __name__ == "__main__":
    unittest.main()
