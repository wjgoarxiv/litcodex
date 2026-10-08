"""Gate checks for decks built from tonality packs (OF-110..OF-114 and pack-aware geometry).

The engine stamps a pack deck in three places the gate reads: the core-properties subject
(`lit-pptx tonality=… density=… grid=… variance=…`), a `family@<id>` name on the first shape of
each slide, and a `title@<treatment>` name on every title frame. These fixtures write the same
marks by hand so each check is exercised without the Node runtime.
"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.dont_write_bytecode = True
sys.path.insert(0, str(ROOT / 'plugins/litcodex/skills/lit-pptx/scripts'))
from craft_extras import analyze_craft, declarative_title  # noqa: E402
import inventory  # noqa: E402
import qa_deck  # noqa: E402
import validate_pptx  # noqa: E402

STAMP = 'lit-pptx tonality=ledger density=10 grid=compact variance=5'


def new_deck(stamp=STAMP):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    if stamp:
        prs.core_properties.subject = stamp
    return prs


def rect(slide, x, y, w, h, fill='0E6B5A', name=None, outline=None):
    item = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    if fill:
        item.fill.solid()
        item.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        item.fill.background()
    if outline:
        item.line.color.rgb = RGBColor.from_string(outline)
    else:
        item.line.fill.background()
    if name:
        item.name = name
    return item


def words(slide, x, y, w, h, value, size=14, name=None, colour='16212C', bold=False):
    frame = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    frame.text_frame.word_wrap = True
    frame.text_frame.text = value
    for paragraph in frame.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(colour)
    if name:
        frame.name = name
    return frame


def page(prs, family, title, treatment, size=26):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    marker = rect(slide, 0, 0, 0, 0, name=f'family@{family}')
    marker.width = marker.height = 0
    geometry = {
        'top-rule': (24, 24, 912, 40), 'side-rail': (24, 36, 288, 90),
        'bottom-anchor': (24, 404, 912, 60), 'section': (24, 160, 760, 60),
    }[treatment]
    words(slide, *geometry, title, size=size, name=f'title@{treatment}', bold=True)
    if treatment == 'top-rule':
        rect(slide, 24, 70, 60, 3)
    return slide


def fill_body(slide, x=24, top=96, bottom=480, columns=1, width=912):
    step = (bottom - top) / 4
    col_w = (width - 24 * (columns - 1)) / columns
    for c in range(columns):
        for r in range(4):
            words(slide, x + c * (col_w + 24), top + r * step, col_w, step - 8, '근거가 되는 사실 한 줄과 그 출처', 14)


def titles_last(prs):
    """The engine draws each title after the body; keep that z-order in the fixtures."""
    for slide in prs.slides:
        tree = slide.shapes._spTree
        for shape in [s for s in slide.shapes if s.name.startswith('title@')]:
            tree.remove(shape._element)
            tree.append(shape._element)
    return prs


def saved(prs):
    titles_last(prs)
    handle = tempfile.NamedTemporaryFile(suffix='.pptx', delete=False)
    handle.close()
    prs.save(handle.name)
    return Path(handle.name)


def found(prs, rule, severity='HIGH'):
    return [f for f in analyze_craft(titles_last(prs))['findings'] if f['rule'] == rule and f['severity'] == severity]


def flat_deck():
    """One title treatment and one short column repeated: the look before tonality packs."""
    prs = new_deck()
    cover = page(prs, 'cover-typographic', '3분기 운영 점검', 'section', size=44)
    del cover
    for n in range(2, 10):
        slide = page(prs, 'text-column', f'점검 항목 {n}', 'top-rule')
        words(slide, 24, 96, 640, 60, '- 근거 한 줄\n- 근거 두 줄', 16)
    title = prs.slides[4].shapes[1]
    title.left = Pt(36)
    prs.slides[6].shapes[1].name = 'title@side-rail'
    return prs


def varied_deck():
    prs = new_deck()
    page(prs, 'cover-typographic', '3분기 운영 점검', 'section', size=44)
    plan = [('text-column', 'top-rule', 1), ('comparison', 'side-rail', 2), ('kpi-row', 'bottom-anchor', 3),
            ('table-insight', 'top-rule', 2), ('timeline', 'side-rail', 1), ('matrix-2x2', 'bottom-anchor', 2),
            ('text-column', 'top-rule', 3), ('process', 'side-rail', 3)]
    for n, (family, treatment, columns) in enumerate(plan, start=2):
        slide = page(prs, family, f'점검 항목 {n}', treatment)
        if treatment == 'side-rail':
            fill_body(slide, x=336, top=36, bottom=480, columns=columns, width=600)
        elif treatment == 'bottom-anchor':
            fill_body(slide, top=36, bottom=392, columns=columns)
        else:
            fill_body(slide, columns=columns)
    return prs


class DeckOutputChecks(unittest.TestCase):
    def test_flat_deck_fails_every_output_check(self):
        prs = flat_deck()
        details = {f['rule']: f for f in found(prs, 'OF-110') + found(prs, 'OF-111') + found(prs, 'OF-112') + found(prs, 'OF-113')}
        self.assertEqual(set(details), {'OF-110', 'OF-111', 'OF-112', 'OF-113'})
        self.assertTrue(any(f['slide'] == 7 and 'named side-rail' in f['detail'] for f in found(prs, 'OF-110')))
        self.assertTrue(any(f['slide'] is None and '1 title treatment' in f['detail'] for f in found(prs, 'OF-110')))
        self.assertTrue(any(f['slide'] == 5 for f in found(prs, 'OF-113')))

    def test_varied_deck_passes_the_output_checks(self):
        prs = varied_deck()
        for rule in ('OF-110', 'OF-111', 'OF-112', 'OF-113'):
            self.assertEqual(found(prs, rule), [], rule)

    def test_a_deck_without_pack_marks_is_not_judged_on_output(self):
        prs = new_deck(stamp=None)
        for _ in range(9):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            words(slide, 24, 24, 900, 40, 'Plain title', 26)
            words(slide, 24, 96, 640, 60, 'One line', 16)
        for rule in ('OF-110', 'OF-111', 'OF-112', 'OF-113', 'OF-114'):
            self.assertEqual(found(prs, rule) + found(prs, rule, 'MEDIUM'), [], rule)


class TitleLabelChecks(unittest.TestCase):
    def test_sentences_and_labels(self):
        for sentence in ('매출은 전년보다 23% 늘었다', '처리량이 목표를 넘었습니다', '이익률 개선됨', 'Revenue grew 23% on volume',
                         'Margins are stable', 'The pilot worked.'):
            self.assertTrue(declarative_title(sentence), sentence)
        for label in ('분기별 매출 추이', '원가 상승 품목별 대응 우선순위', '3분기 계획 대비 실적 (예시)', 'Conversion by feed rate',
                      'Why did returns rise?', '“Speed is a feature”'):
            self.assertFalse(declarative_title(label), label)

    def test_sentence_title_fails_under_a_pack(self):
        prs = varied_deck()
        prs.slides[3].shapes[1].text_frame.text = '매출은 전년보다 23% 늘었다'
        hits = found(prs, 'OF-114')
        self.assertEqual([f['slide'] for f in hits], [4])
        self.assertIn('label', hits[0]['hint'])


class PackSlideGeometry(unittest.TestCase):
    def test_pack_section_slide_is_a_display_slide(self):
        prs = new_deck()
        slide = page(prs, 'section-rule', '원가 상승 이슈와 대응', 'section', size=44)
        words(slide, 24, 260, 912, 24, '원가 상승 품목별 대응 우선순위', 18)
        words(slide, 24, 290, 912, 40, '즉시 대응하는 두 항목은 11월까지 끝낸다', 13)
        self.assertEqual(found(prs, 'OF-109'), [])

    def test_a_field_bleeding_off_the_canvas_is_not_a_card(self):
        prs = new_deck()
        slide = page(prs, 'comparison', '계획 대비 실적 비교', 'side-rail')
        field = rect(slide, 0, 0, 324, 540, fill='F2F4F6')
        slide.shapes._spTree.remove(field._element)
        slide.shapes._spTree.insert(2, field._element)
        fill_body(slide, x=336, top=36, bottom=480, columns=2, width=600)
        self.assertEqual([f for f in found(prs, 'OF-109') if 'card-row' in f['detail']], [])

    def test_side_rail_body_is_measured_from_the_top(self):
        prs = new_deck()
        slide = page(prs, 'text-column', '계획 대비 실적', 'side-rail')
        fill_body(slide, x=336, top=36, bottom=480, columns=1, width=600)
        self.assertEqual(found(prs, 'OF-109'), [])

    def test_a_chart_is_content(self):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        prs = new_deck()
        slide = page(prs, 'chart-insight', '지역별 매출 증가', 'top-rule')
        data = CategoryChartData()
        data.categories = ['북부', '서부']
        data.add_series('증가', [710, 290])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(24), Pt(84), Pt(560), Pt(390), data)
        words(slide, 610, 84, 326, 60, '북부와 서부가 증가분의 71%를 차지한다', 13)
        self.assertEqual(found(prs, 'OF-109'), [])

    def test_korean_lines_wrap_at_korean_widths(self):
        prs = new_deck()
        slide = page(prs, 'text-column', '원가 대응 일정과 담당', 'top-rule')
        text = ('즉시 대응하는 두 항목(포장재, 물류)은 11월까지 끝낸다 (표 2)\n'
                '원자재 장기 계약은 12월에 안을 확정하고 2월부터 고정가를 적용한다\n'
                '보류한 설비 교체(연 2억 원)는 1월 점검 결과를 보고 3월에 다시 검토한다')
        words(slide, 24, 96, 385, 92, text, 13)
        words(slide, 24, 300, 900, 60, '한 줄로 길게 이어지는 문장이 화면 너비를 거의 다 쓰면서 마흔 글자를 넘기는 경우\n'
              '두 번째 줄도 마찬가지로 길게 이어져서 읽는 눈이 다음 줄의 시작을 놓치게 된다', 13)
        hits = found(prs, 'OF-102')
        self.assertEqual(len(hits), 1, hits)
        self.assertGreater(float(hits[0]['detail'].split()[0]), 38)

    def test_the_notice_line_does_not_fill_an_empty_body(self):
        prs = new_deck()
        slide = page(prs, 'text-column', '계획 대비 실적', 'top-rule')
        words(slide, 24, 96, 600, 40, '한 줄뿐인 본문', 16)
        words(slide, 24, 497, 160, 18, '예시 데이터 — 실제 수치로 바꿔 주세요', 9, name='lit-notice text')
        self.assertTrue(any('largest empty band' in f['detail'] for f in found(prs, 'OF-109')))


class PackGateCompanions(unittest.TestCase):
    def test_a_treatment_title_counts_wherever_it_sits(self):
        prs = new_deck()
        slide = page(prs, 'section-rule', '원가 상승 이슈와 대응', 'section', size=44)
        for y in (260, 300, 340):
            words(slide, 24, y, 900, 30, '원가 상승 품목별 대응 우선순위', 14)
        words(slide, 24, 470, 600, 16, '출처: 예시 자료', 9)
        report = validate_pptx.lint(saved(prs))
        self.assertTrue(report['checks']['missing_titles']['pass'], report['checks']['missing_titles'])

    def test_a_closing_references_page_is_the_source_note(self):
        prs = new_deck()
        slide = page(prs, 'references-appendix', '참고 자료', 'top-rule')
        words(slide, 24, 96, 900, 30, 'References', 18)
        words(slide, 24, 130, 900, 300, '[1] 예시 저자, 예시 보고서, 2026.', 13)
        self.assertTrue(validate_pptx.lint(saved(prs))['checks']['missing_source_note']['pass'])

    def test_text_on_a_dark_slide_ground_is_judged_against_that_ground(self):
        def deck(colour):
            prs = new_deck()
            slide = page(prs, 'text-column', '야간 지표', 'top-rule')
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string('0F1720')
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.color.rgb = RGBColor.from_string(colour)
            return qa_deck.check_contrast(saved(prs))
        self.assertTrue(deck('E8EDF2')['pass'])
        self.assertFalse(deck('1B2430')['pass'])

    def test_latin_text_is_measured_at_latin_widths(self):
        prs = new_deck()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        fits = words(slide, 24, 24, 400, 20, 'Revenue by region, third quarter of 2026', 14)
        tight = words(slide, 24, 100, 120, 18, 'A sentence that cannot possibly fit into this small frame at all', 14)
        broken = words(slide, 24, 200, 400, 20, 'first line\vsecond line\vthird line', 14)
        del fits, tight, broken
        issues = inventory.inspect(saved(prs)).get('slide-1', {})
        self.assertNotIn('shape-1', issues)
        self.assertIn('frame', issues['shape-2']['overflow'])
        self.assertIn('frame', issues['shape-3']['overflow'])

    def test_pack_cover_is_not_judged_for_fill_and_body_counts_from_the_title(self):
        prs = new_deck()
        cover = page(prs, 'cover-typographic', '3분기 운영 점검', 'section', size=44)
        words(cover, 24, 30, 400, 20, '예시 기업 경영기획팀', 12)
        slide = page(prs, 'text-column', '계획 대비 실적', 'top-rule')
        for y in (84, 190, 296, 402):
            words(slide, 24, y, 912, 96, '근거가 되는 사실과 그 출처를 담은 본문 단락', 14)
        result = qa_deck.check_content_quality(saved(prs))
        self.assertEqual([v for v in result['violations'] if v['kind'] == 'underfilled'], [], result)

    def test_pack_body_is_measured_by_the_columns_it_occupies(self):
        def deck(rows):
            prs = new_deck()
            slide = page(prs, 'text-column', '시범 운영 검토 배경과 결과', 'top-rule')
            for x in (24, 492):
                for r in range(rows):
                    words(slide, x, 96 + r * 96, 444, 16, '항목 제목', 12)
                    words(slide, x, 116 + r * 96, 444, 24, '(1) 근거가 되는 사실 한 줄', 12)
            return qa_deck.check_content_quality(saved(prs))
        self.assertEqual([v for v in deck(4)['violations'] if v['kind'] == 'underfilled'], [], deck(4))
        self.assertTrue([v for v in deck(1)['violations'] if v['kind'] == 'underfilled'])

    def test_an_outline_around_content_is_its_frame(self):
        prs = new_deck()
        slide = page(prs, 'matrix-2x2', '대응 우선순위', 'top-rule')
        for x, y in ((24, 96), (492, 96), (24, 290), (492, 290)):
            rect(slide, x, y, 444, 180, fill=None, outline='D0D6DC')
            words(slide, x + 12, y + 12, 420, 150, '항목과 근거', 14)
        rect(slide, 400, 20, 30, 30, fill=None, outline='000000')
        result = qa_deck.check_content_quality(saved(prs))
        self.assertEqual(len([v for v in result['violations'] if v['kind'] == 'empty-shape']), 1, result)


if __name__ == '__main__':
    unittest.main()
