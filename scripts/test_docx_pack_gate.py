"""Document gate checks for tonality documents (docx_craft.py as the lit-docx gate).

Run with the office runner's docx Python (python-docx, PyYAML, Markdown, PyMuPDF); scripts/office-engine.test.mjs
does that and fails if this file is skipped.
"""
from __future__ import annotations

import json
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import unittest
import zlib
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / 'plugins/litcodex/skills/lit-docx'
sys.dont_write_bytecode = True
sys.path.insert(0, str(SKILL / 'scripts'))
from docx import Document  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

import docx_craft  # noqa: E402

SOURCE = """---
title: 3분기 물류 운영 점검
tonality: Report
notice: "예시 데이터: 실제 수치로 바꿔 주세요"
---

## 점검 배경

창고 두 곳의 출고 처리량과 지연 원인을 분기 단위로 정리한다.

## 처리량 추이

3분기 출고는 41만 건으로 2분기보다 6% 많았다.
"""


MEMO = """---
title: New loading-dock hours from 3 November
tonality: Memo
author: Site operations (example)
organization: Example Logistics
date: 2026-10-06
short_title: Loading-dock hours
notice: "Example memo: names, dates and numbers are placeholders"
---

::: cover variant=masthead kicker="Internal memo"
To: Dispatch and carrier desks (example) · From: Site operations · Subject: loading-dock hours
:::

# What changes

From Monday 3 November, the north loading docks open at 05:30 instead of 07:00 and close at 21:00. Inbound trailers book a 30-minute slot in the dock system; a trailer without a slot waits in the holding yard, not on the access road.

We are making this change because inbound trailers now arrive about 48 times a day, up from 30 a year ago. Trailers waiting on the access road have blocked the bus stop three times since August.

::: callout kind=warning title="Access road"
The access road in front of the gatehouse is a bus route. Trailers never wait there, even with the engine off. Security will send any waiting trailer to the holding yard.
:::

# What you need to do

1. Book every inbound trailer in the dock system at least one working day ahead. The system issues a slot code with the booking email.
2. Tell your carrier to follow the "Holding yard" signs from the roundabout when they arrive early.
3. Meet early drivers at the holding-yard gate. The gatehouse still checks papers, but drivers walk to it through the yard.

Refrigerated trailers keep using the two powered bays beside the gatehouse. Nothing changes for them.

# Numbers behind the decision

Table 1. Dock use, one year apart (example data)

| Measure (example) | Oct 2025 | Sep 2026 |
| --- | ---: | ---: |
| Inbound trailers per day | 30 | 48 |
| Average waiting time (min) | 22 | 41 |
| Access-road blockages reported | 0 | 3 |

# Questions

Ask the site operations desk on extension 2210 (example) or reply to this memo. We will review the new hours at the end of January.
"""
SPILL = """
The new hours apply to parcel vans as well; their drivers receive the same slot code with their route sheet and wait at the west entrance when the docks are full.

We will put up the new signs in the week before 3 November and take down the old dock-hour boards at the gatehouse on the same day.

The night shift keeps its present break times; the extra early hour is covered by two staff moved from the late shift, who finish at 14:00 instead of 15:30.
"""


def convert(source: str, folder: Path, name: str = 'doc') -> Path:
    md = folder / f'{name}.md'
    out = folder / f'{name}.docx'
    md.write_text(source, encoding='utf-8')
    result = subprocess.run([sys.executable, str(SKILL / 'scripts/convert_md_to_docx.py'), str(md), str(out)],
                            capture_output=True, text=True)
    if result.returncode:
        raise AssertionError(result.stderr or result.stdout)
    return out


def gate(docx: Path, *flags: str) -> tuple[int, dict]:
    result = subprocess.run([sys.executable, str(SKILL / 'scripts/docx_craft.py'), str(docx), *flags],
                            capture_output=True, text=True)
    return result.returncode, json.loads(result.stdout)


def flat_png(w: int, h: int) -> bytes:
    """A flat grey PNG (the docx runtime has no imaging library)."""
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data))
    rows = b''.join(b'\x00' + bytes((90, 110, 130)) * w for _ in range(h))
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(rows)) + chunk(b'IEND', b''))


def document_xml(docx: Path) -> str:
    import zipfile
    with zipfile.ZipFile(docx) as z:
        return z.read('word/document.xml').decode('utf-8')


def ppr(xml: str, start: str) -> str:
    """The paragraph properties of the first paragraph whose text starts with `start`."""
    for para in re.findall(r'<w:p[ >].*?</w:p>', xml, flags=re.S):
        text = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', para)).replace('\xa0', ' ')
        if text.startswith(start):
            found = re.search(r'<w:pPr>.*?</w:pPr>', para, flags=re.S)
            return found.group(0) if found else ''
    raise AssertionError(f'no paragraph starts with {start!r}')


def spacing(props: str, side: str) -> int:
    found = re.search(rf'<w:spacing[^>]*w:{side}="(\d+)"', props)
    return int(found.group(1)) if found else 0


def keeps_next(props: str) -> bool:
    return '<w:keepNext/>' in props


class DocumentGate(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_a_restrained_tonality_document_passes(self):
        docx = convert(SOURCE, self.dir)
        code, report = gate(docx, '--source', str(self.dir / 'doc.md'))
        self.assertEqual(code, 0, report)
        self.assertTrue(report['pass'])
        self.assertEqual(report['output']['findings'], [])

    def test_a_sentence_heading_fails(self):
        docx = convert(SOURCE.replace('## 처리량 추이', '## 처리량은 6% 늘었다'), self.dir)
        code, report = gate(docx, '--source', str(self.dir / 'doc.md'))
        self.assertEqual(code, 1)
        self.assertIn('heading.declarative', {f['check'] for f in report['output']['findings']})

    def test_a_coloured_heading_fails_under_a_tonality(self):
        docx = convert(SOURCE, self.dir)
        document = Document(str(docx))
        document.styles['Heading 2'].font.color.rgb = RGBColor(0x12, 0x3A, 0x6B)
        document.styles['Heading 2'].font.size = Pt(20)
        document.save(str(docx))
        code, report = gate(docx, '--source', str(self.dir / 'doc.md'))
        self.assertEqual(code, 1)
        self.assertIn('heading.ink', {f['check'] for f in report['output']['findings']})

    def test_an_iso_date_in_korean_text_fails_on_every_path(self):
        source = SOURCE.replace('창고 두 곳의', '2026-09-30 기준 창고 두 곳의').replace('tonality: Report\n', '')
        md = self.dir / 'plain.md'
        md.write_text(source, encoding='utf-8')
        docx = self.dir / 'plain.docx'
        document = Document()
        document.add_paragraph('2026-09-30 기준 창고 두 곳의 출고 처리량')
        document.save(str(docx))
        code, report = gate(docx, '--source', str(md))
        self.assertEqual(code, 1)
        self.assertIn('date.iso', {f['check'] for f in report['output']['findings']})

    def test_a_missing_heading_is_a_content_failure(self):
        docx = convert(SOURCE, self.dir)
        md = self.dir / 'other.md'
        md.write_text(SOURCE + '\n## 다음 분기 계획\n\n계획 한 줄\n', encoding='utf-8')
        code, report = gate(docx, '--source', str(md))
        self.assertEqual(code, 1)
        self.assertEqual(report['content']['missing_headings'], ['다음 분기 계획'])

    def test_one_document_compared_with_itself_fails_the_structure_check(self):
        report_doc = convert(SOURCE, self.dir, 'report')
        same = convert(SOURCE, self.dir, 'same')
        code, report = gate(report_doc, '--source', str(self.dir / 'report.md'), '--compare', str(same))
        self.assertEqual(code, 1)
        self.assertIn('tonality.structure', {f['check'] for f in report['output']['findings']})

    @unittest.skipUnless(shutil.which('soffice'), 'soffice renders the pages')
    def test_a_memo_that_fits_one_page_stays_on_it_and_a_spill_is_reported(self):
        docx = convert(MEMO, self.dir, 'memo')
        paragraphs = Document(str(docx)).paragraphs
        head = next(p for p in paragraphs if p.text.strip() == 'Access road')
        after = next(p for p in paragraphs if p.text.startswith('Refrigerated trailers'))
        self.assertGreaterEqual(head.paragraph_format.space_before, Pt(12), 'a run-in head opens its block')
        self.assertGreaterEqual(after.paragraph_format.space_before or 0, Pt(6), 'a paragraph after a list stands clear of it')
        code, report = gate(docx, '--source', str(self.dir / 'memo.md'), '--layout', '--tonality', 'Memo')
        self.assertTrue(report['output']['rendered'], report['output'].get('render'))
        self.assertEqual(report['output']['pages'], 1, report['output'])
        self.assertNotIn('memo.fit', {f['check'] for f in report['output']['findings']})
        long = convert(MEMO + SPILL, self.dir, 'long')
        code, report = gate(long, '--source', str(self.dir / 'long.md'), '--layout', '--tonality', 'Memo')
        self.assertEqual(code, 1)
        self.assertIn('memo.fit', {f['check'] for f in report['output']['findings']})

    def test_a_publisher_profile_keeps_text_clear_of_tables_and_captions(self):
        (self.dir / 'fig.png').write_bytes(flat_png(60, 30))
        md = self.dir / 'paper.md'
        md.write_text(
            '---\ntitle: Throughput of a split picking wave\nauthor: Example Author A\n---\n\n# Results\n\n'
            'The split wave picks more orders than the single wave at every staffing level.\n\n'
            'Table 1. Orders picked per hour (example values)\n\n| Pickers | Single wave | Split wave |\n| ---: | ---: | ---: |\n'
            '| 8 | 410 | 520 |\n| 12 | 560 | 735 |\n\nAcross both staffing levels the split wave picks about a quarter more.\n\n'
            '![Figure 1. Orders picked over a shift (synthetic)](fig.png)\n\nThe two curves differ in shape as well as height.\n',
            encoding='utf-8')
        for publisher in ('elsevier', 'acs'):
            out = self.dir / f'{publisher}.docx'
            result = subprocess.run([sys.executable, str(SKILL / 'scripts/convert_md_to_docx.py'), str(md), str(out), '--publisher', publisher],
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            paragraphs = Document(str(out)).paragraphs
            find = lambda start: next(p for p in paragraphs if p.text.replace('\xa0', ' ').startswith(start))  # noqa: E731
            self.assertGreaterEqual(find('Across both').paragraph_format.space_before or 0, Pt(8), f'{publisher}: the paragraph after the table')
            self.assertGreaterEqual(find('Table 1').paragraph_format.space_before or 0, Pt(8), f'{publisher}: the table caption after text')
            self.assertGreaterEqual(find('Figure 1').paragraph_format.space_after or 0, Pt(8), f'{publisher}: the figure caption keeps 8 pt under it')

    def test_a_short_list_keeps_together_and_a_heading_keeps_its_lead_with_the_table(self):
        source = ('---\ntitle: 자동 분류 설비 운영 결과\ntonality: Report\nnotice: "예시 데이터: 수치는 가정입니다"\n---\n\n'
                  '# 결론과 다음 단계\n\n자동 분류 설비는 처리량과 정확도에서 기대한 효과를 냈다.\n\n다음 단계는 아래와 같다.\n\n'
                  '1. 2026년 11월: 투입부 개조 예산 승인\n2. 2027년 1월: 주말 두 차례에 걸쳐 개조 공사\n'
                  '3. 2027년 2월: 야간 정비 인력 1명 추가 배치\n4. 2027년 3월: 개조 효과 확인 뒤 상세 설계 착수\n\n'
                  '이 보고서의 모든 수치는 예시다.\n\n# 시나리오 비교\n\n세 시나리오의 투자비와 회수 기간을 비교했다.\n\n'
                  '표 1. 확대 시나리오 비교 (예시)\n\n| 시나리오 | 투자비 (억 원) | 회수 기간 (년) |\n| --- | ---: | ---: |\n'
                  '| A. 현재 설비 | 11.8 | 6.2 |\n| B. 전용 투입구 | 12.9 | 5.0 |\n\n자료: 예시 추정\n')
        xml = document_xml(convert(source, self.dir))
        self.assertTrue(keeps_next(ppr(xml, '다음 단계는 아래와 같다.')), 'the lead-in keeps with its list')
        for item in ('2026년 11월', '2027년 1월', '2027년 2월'):
            self.assertTrue(keeps_next(ppr(xml, item)), f'{item} keeps with the next item')
        self.assertFalse(keeps_next(ppr(xml, '2027년 3월')), 'the last item ends the chain')
        self.assertTrue(keeps_next(ppr(xml, '세 시나리오의 투자비와')), "the heading's lead keeps with the table it opens")

    def test_the_page_checks_find_a_split_list_a_heading_apart_a_near_empty_second_page_and_a_lone_last_page(self):
        import docx_layout as L
        frame = {'page_h': 841.9, 'top': 70.0, 'bottom': 70.0}

        def page(lines, bottom):
            return {'h': 841.9, 'lines': lines, 'ink': [(y0, y1) for y0, y1, _ in lines] + [(70.0, bottom)], 'images': 0,
                    'boxes': [(72.0, y0, 300.0, y1, t) for y0, y1, t in lines], 'filled': 0.0}
        blocks = [{'t': 'h', 'level': 1, 'text': '결론', 'raw': '결론'}, {'t': 'p', 'text': '다음단계는아래와같다', 'list': False},
                  {'t': 'p', 'text': '첫째항목입니다', 'list': True}, {'t': 'p', 'text': '둘째항목입니다', 'list': True},
                  {'t': 'p', 'text': '셋째항목입니다', 'list': True}, {'t': 'h', 'level': 1, 'text': '비교', 'raw': '비교'},
                  {'t': 'p', 'text': '세시나리오를비교했다', 'list': False},
                  {'t': 'tbl', 'rows': ['시나리오투자비', 'A현재설비'], 'component': None, 'cells': []}]
        info = {'blocks': blocks, 'components': {k: 0 for k in L.KINDS}, 'cover': False, 'memo': False, 'frame': frame}
        p1 = page([(80, 92, '결론'), (100, 112, '다음단계는아래와같다'), (700, 712, '첫째항목입니다')], 712)
        p2 = page([(80, 92, '둘째항목입니다'), (100, 112, '셋째항목입니다'), (740, 752, '비교'), (760, 770, '세시나리오를비교했다')], 770)
        p3 = page([(80, 90, '시나리오투자비'), (95, 105, 'A현재설비')], 105)
        checks = {f['check'] for f in L.paged(info, [p1, p2, p3], True, {})[0]}
        for check in ('list.split', 'heading.apart', 'page.spill'):
            self.assertIn(check, checks)
        lone = page([(80, 92, '끝')], 92)
        self.assertIn('page.spill', {f['check'] for f in L.paged({**info, 'blocks': blocks[:1]}, [p1, p1, lone], True, {})[0]})

    def test_a_heading_between_a_lead_and_a_table_owns_the_table(self):
        import docx_layout as L
        frame = {'page_h': 841.9, 'top': 70.0, 'bottom': 70.0}

        def page(lines, bottom):
            return {'h': 841.9, 'lines': lines, 'ink': [(y0, y1) for y0, y1, _ in lines] + [(70.0, bottom)], 'images': 0,
                    'boxes': [(72.0, y0, 300.0, y1, t) for y0, y1, t in lines], 'filled': 0.0}
        # A short section ends page 1; the next heading, its lead and its own table open page 2.
        # The body lines are drawn, not read as blocks, so the second heading follows the first lead in block order.
        body = [f'본문{k:02d}줄입니다' for k in range(30)]
        blocks = [{'t': 'h', 'level': 1, 'text': '배경', 'raw': '배경'}, {'t': 'p', 'text': '두창고를비교했다', 'list': False},
                  {'t': 'h', 'level': 1, 'text': '비용', 'raw': '비용'}, {'t': 'p', 'text': '안별비용은아래와같다', 'list': False},
                  {'t': 'tbl', 'rows': ['안투자비', 'A현재설비'], 'component': None, 'cells': []}]
        info = {'blocks': blocks, 'components': {k: 0 for k in L.KINDS}, 'cover': False, 'memo': False, 'frame': frame}
        p1 = page([(80, 92, '배경'), (100, 112, '두창고를비교했다')] + [(120 + 20 * k, 132 + 20 * k, t) for k, t in enumerate(body)], 712)
        p2 = page([(80, 92, '비용'), (100, 112, '안별비용은아래와같다'), (120, 130, '안투자비'), (135, 145, 'A현재설비')], 145)
        checks = {f['check'] for f in L.paged(info, [p1, p2], True, {})[0]}
        self.assertNotIn('heading.apart', checks, 'a heading between the lead and the table owns the table')

    @unittest.skipUnless(shutil.which('soffice'), 'soffice renders the pages')
    def test_a_source_of_about_a_page_is_set_tight_and_a_section_after_a_ruled_box_draws_no_second_rule(self):
        source = (MEMO.replace('tonality: Memo', 'tonality: Brief')
                  .replace('3. Meet early drivers at the holding-yard gate. The gatehouse still checks papers, but drivers walk to it through the yard.',
                           '3. Meet early drivers at the holding-yard gate; the gatehouse still checks papers.')
                  .replace('Refrigerated trailers keep using the two powered bays beside the gatehouse. Nothing changes for them.',
                           'Refrigerated trailers keep the two powered bays beside the gatehouse.'))
        md = self.dir / 'brief.md'
        md.write_text(source, encoding='utf-8')
        out = self.dir / 'brief.docx'
        result = subprocess.run([sys.executable, str(SKILL / 'scripts/convert_md_to_docx.py'), str(md), str(out)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('tight spacing', result.stdout)
        self.assertIn('<w:top w:val="nil"/>', ppr(document_xml(out), 'What you need to do'), 'no hairline over the section under the box')
        code, report = gate(out, '--source', str(md), '--layout', '--tonality', 'Brief')
        self.assertTrue(report['output']['rendered'], report['output'].get('render'))
        self.assertEqual(report['output']['pages'], 1, report['output'])
        self.assertNotIn('page.spill', {f['check'] for f in report['output']['findings']})

    def test_a_wrapping_column_takes_room_negatives_take_the_minus_sign_and_an_uncaptioned_table_stands_clear(self):
        source = ('---\ntitle: 센터별 주간 입고\ntonality: Report\nnotice: "예시 데이터: 수치는 가정입니다"\n---\n\n# 센터별 입고\n\n'
                  '센터마다 입고와 변화를 적었다.\n\n| Site | Trailers | Change | Share |\n| --- | ---: | ---: | ---: |\n'
                  '| Harbour Road | 1,240 | -3 | 21% |\n| North Gate Cross-Dock | 980 | -12 | 17% |\n| Old Mill | 2,105 | 8 | 36% |\n\n자료: 예시 집계\n')
        xml = document_xml(convert(source, self.dir))
        grid = [int(w) / 20 for w in re.findall(r'<w:gridCol w:w="(\d+)"/>', xml)]
        self.assertGreaterEqual(grid[0], 95, f'the site column holds its longest name or most of it: {grid}')
        self.assertIn('\u22123<', xml)
        self.assertIn('\u221212<', xml)
        self.assertNotIn('>-3<', xml)
        self.assertGreaterEqual(spacing(ppr(xml, '센터마다 입고와'), 'after'), 120, 'the sentence over an uncaptioned table keeps 6 pt under it')

    def test_a_publisher_profile_spaces_lists_and_tables_and_keeps_a_short_list_together(self):
        md = self.dir / 'p.md'
        md.write_text('---\ntitle: A five-line shift report\nauthor: Example Author A\n---\n\n# Results\n\n'
                      'We propose that every short shift report states five lines.\n\n1. Orders picked at the end of the shift.\n'
                      '2. Half-wave time.\n3. Peak dock load.\n\nLines 3 and 4 together let a reader compute the margin.\n\n'
                      'Table 1. Margin (example)\n\n| Inflow (lines/min) | Margin (lines) |\n| ---: | ---: |\n| 180 | 0.88 |\n| 220 | -0.07 |\n\n'
                      'The negative margin at 220 lines a minute does not make the layout unusable.\n', encoding='utf-8')
        for publisher in ('elsevier', 'acs'):
            out = self.dir / f'{publisher}.docx'
            result = subprocess.run([sys.executable, str(SKILL / 'scripts/convert_md_to_docx.py'), str(md), str(out), '--publisher', publisher],
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            xml = document_xml(out)
            self.assertGreaterEqual(spacing(ppr(xml, 'Lines 3 and 4'), 'before'), 120, f'{publisher}: after the list')
            self.assertGreaterEqual(spacing(ppr(xml, 'The negative margin'), 'before'), 180, f'{publisher}: after the table')
            self.assertTrue(keeps_next(ppr(xml, 'Orders picked at')) and keeps_next(ppr(xml, 'Half-wave time')), f'{publisher}: the list keeps together')

    def test_a_notice_joined_by_a_spaced_dash_fails_under_a_tonality(self):
        dashed = SOURCE.replace('notice: "예시 데이터: 실제 수치로 바꿔 주세요"', 'notice: 예시 안내문 — 일정은 가정입니다')
        docx = convert(dashed, self.dir)
        code, report = gate(docx, '--source', str(self.dir / 'doc.md'))
        self.assertEqual(code, 1)
        self.assertIn('notice.dash', {f['check'] for f in report['output']['findings']})
        docx = convert(SOURCE, self.dir, 'colon')
        code, report = gate(docx, '--source', str(self.dir / 'colon.md'))
        self.assertNotIn('notice.dash', {f['check'] for f in report['output']['findings']})

    def test_a_two_column_body_spans_a_table_source_lets_a_column_list_break_once_and_balances_its_last_page(self):
        source = ('---\ntitle: Dwell time at two cross-dock yards\nauthor: Example Author A\ntonality: Journal\n'
                  'notice: "Synthetic example: data are placeholders"\n---\n\n# Method\n\nEach yard logged trailers for six weeks.\n\n'
                  'Table 1. Trailers by yard and shift (example values)\n\n| Yard | Shift | Trailers | Dwell (h) |\n| --- | --- | ---: | ---: |\n'
                  '| North | Day | 1,240 | 3.1 |\n| South | Night | 980 | 4.6 |\n\nSource: synthetic data, six weeks\n\n## Logging\n\n'
                  'Each trailer was timed from gate-in to gate-out.\n\nEvery short yard report states five lines:\n\n'
                  '1. Trailers per shift.\n2. Median dwell time.\n3. Longest wait at the gate.\n4. Doors staffed per shift.\n5. Number of weeks logged.\n\n'
                  'Lines 3 and 4 together give the door load.\n')
        xml = document_xml(convert(source, self.dir))
        table = xml.index('Trailers by yard and shift')
        self.assertLess(xml.index('>Source: synthetic data'), xml.index('w:num="1"', table), 'the source line stands under the page-wide table')
        for start in ('Every short yard report', 'Trailers per shift', 'Doors staffed per shift'):
            self.assertTrue(keeps_next(ppr(xml, start)), f'{start} keeps with the next paragraph')
        for start in ('Median dwell time', 'Longest wait at the gate'):
            self.assertFalse(keeps_next(ppr(xml, start)), f'the column list may break after {start}')
        sects = list(re.finditer(r'<w:sectPr\b.*?</w:sectPr>', xml, flags=re.S))
        self.assertNotIn('w:num="2"', sects[-1].group(0), 'the document closes on a one-column section')
        self.assertIn('w:num="2"', sects[-2].group(0))
        self.assertGreater(sects[-2].start(), xml.index('Lines 3 and 4 together'))
        self.assertIn('<w:widowControl w:val="0"/>', ppr(xml, 'Lines 3 and 4 together'), 'the last paragraph may break across the balanced columns')

    def test_a_short_closing_section_keeps_whole_and_one_with_a_table_does_not(self):
        head = ('---\ntitle: 자동 분류 설비 운영 결과\ntonality: Report\nnotice: "예시 데이터: 수치는 가정입니다"\n---\n\n# 배경\n\n'
                '자동 분류 설비를 제2물류센터에서 1년 운영했고, 처리량과 정확도, 비용 회수를 분기마다 기록해 계획과 비교했다.\n\n')
        tail = ('# 결론과 다음 단계\n\n자동 분류 설비는 처리량과 정확도에서 기대한 효과를 냈다. 비용 회수가 늦어진 원인은 설비 자체가 아니라 '
                '투입부 설계와 야간 운영에 있었고, 두 가지 모두 고칠 수 있다.\n\n다음 단계는 아래와 같다.\n\n1. 2026년 11월: 투입부 개조 예산 승인\n'
                '2. 2027년 1월: 개조 공사\n3. 2027년 2월: 야간 정비 인력 추가 배치\n\n이 보고서의 모든 수치는 예시다.\n')
        xml = document_xml(convert(head + tail, self.dir))
        self.assertTrue(keeps_next(ppr(xml, '자동 분류 설비는 처리량과')), 'the closing paragraph keeps with its list')
        self.assertFalse(keeps_next(ppr(xml, '자동 분류 설비를 제2물류센터에서')), 'an earlier section is untouched')
        xml = document_xml(convert(head + tail + '\n| 단계 | 시기 |\n| --- | --- |\n| 승인 | 11월 |\n', self.dir, 'tabled'))
        self.assertFalse(keeps_next(ppr(xml, '자동 분류 설비는 처리량과')), 'a closing section with a table moves no paragraph')

    def test_the_page_checks_find_a_heading_at_a_column_foot_and_columns_that_do_not_balance(self):
        import docx_layout as L
        frame = {'page_h': 841.9, 'top': 70.0, 'bottom': 70.0}

        def page(lines):
            return {'h': 841.9, 'w': 595.3, 'lines': sorted((y0, y1, t) for _, y0, y1, t in lines), 'ink': [(y0, y1) for _, y0, y1, _ in lines],
                    'images': 0, 'boxes': [(x0, y0, x0 + 230.0, y1, t) for x0, y0, y1, t in lines], 'filled': 0.0}

        def col(x, y0, y1, tag):
            return [(x, y, y + 11, f'{tag}{k}줄') for k, y in enumerate(range(int(y0), int(y1), 14))]

        def blocks(*items):
            return [dict({'t': t, 'text': x, 'list': lst, 'cols': 2}, **({'level': 1, 'raw': x} if t == 'h' else {})) for t, x, lst in items]

        def checks(b, pages):
            info = {'blocks': b, 'components': {k: 0 for k in L.KINDS}, 'cover': False, 'memo': False, 'frame': frame}
            return {f['check'] for f in L.paged(info, pages, True, {})[0]}
        heads = blocks(('h', '배경', False), ('p', '왼쪽0줄', False), ('h', '방법', False), ('p', '다음0줄', False))
        first = page([(72, 80, 92, '배경')] + col(72, 102, 760, '왼쪽') + col(310, 80, 740, '오른쪽') + [(310, 750, 762, '방법')])
        self.assertIn('heading.column', checks(heads, [first, page(col(72, 80, 420, '다음') + col(310, 80, 410, '끝'))]))
        self.assertIn('columns.balance', checks(heads[:2], [first, page(col(72, 80, 500, '다음'))]))
        listed = blocks(('h', '배경', False), ('p', '왼쪽0줄', False), *[('p', f'{n}항목', True) for n in '가나다라마'])
        p1 = page([(72, 80, 92, '배경')] + col(72, 102, 730, '왼쪽') + [(72, 740, 751, '가항목'), (72, 754, 765, '나항목')] + col(310, 80, 765, '오른쪽'))
        p2 = page([(72, 80, 91, '다항목'), (72, 94, 105, '라항목'), (72, 108, 119, '마항목')] + col(72, 124, 420, '다음') + col(310, 80, 420, '끝'))
        self.assertFalse(checks(listed, [p1, p2]) & {'columns.balance', 'list.split', 'heading.column'})
        short = page([(72, 80, 92, '배경')] + col(72, 100, 520, '왼쪽') + col(310, 80, 765, '오른쪽'))
        self.assertIn('columns.balance', checks(listed[:2], [short, p2]))

    def test_the_office_floor_still_runs(self):
        self.assertTrue(callable(docx_craft.audit_docx_office_floor))


if __name__ == '__main__':
    unittest.main()
