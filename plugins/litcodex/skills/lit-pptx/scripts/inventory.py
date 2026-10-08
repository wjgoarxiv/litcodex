#!/usr/bin/env python3
"""Inspect slide geometry and conservative text fit without rendering."""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_POINT = 12700


def box(shape) -> tuple[int, int, int, int]:
    """Return left, top, right, bottom in EMU."""
    return (shape.left, shape.top, shape.left + shape.width, shape.top + shape.height)


def intersection(a, b) -> int:
    width = min(a[2], b[2]) - max(a[0], b[0])
    height = min(a[3], b[3]) - max(a[1], b[1])
    return max(0, width) * max(0, height)


def em_width(text: str) -> float:
    """Advance of a line in ems, a few per cent wider than Pretendard Bold measures.

    Measured Bold advances: Hangul 0.86, em dash 0.97, percent 0.96, capitals about 0.70,
    digits about 0.63, lower case about 0.54, narrow punctuation 0.28, space 0.23. Counting
    every character as a wide glyph made a Latin line read twice its length.
    """
    total = 0.0
    for ch in text:
        if ch == " ":
            total += 0.25
        elif "\uac00" <= ch <= "\ud7a3" or "\u3130" <= ch <= "\u318f":
            total += 0.90
        elif ord(ch) >= 0x1100 or ch == "%":
            total += 0.97
        elif ch in ",.;:'·|!":
            total += 0.30
        elif ch.isupper():
            total += 0.72
        elif ch.isdigit():
            total += 0.66
        elif ch.isalpha():
            total += 0.56
        else:
            total += 0.50
    return total


def text_overflow(shape) -> bool:
    """Flag only clear frame overflow; variable glyph metrics need rendered review."""
    if not shape.has_text_frame or not shape.text.strip():
        return False
    frame = shape.text_frame
    available_width = max(1, shape.width - frame.margin_left - frame.margin_right)
    available_height = max(1, shape.height - frame.margin_top - frame.margin_bottom)
    estimated_lines = 0
    largest_font = 12.0
    for paragraph in frame.paragraphs:
        fonts = [run.font.size.pt for run in paragraph.runs if run.font.size]
        font_size = max(fonts, default=12.0)
        largest_font = max(largest_font, font_size)
        line_width = available_width / (font_size * EMU_PER_POINT)
        # A line break inside the paragraph (\v) starts a new line of its own.
        for line in paragraph.text.split("\v"):
            estimated_lines += max(1, math.ceil(em_width(line) / max(1e-6, line_width)))
    return estimated_lines * largest_font * EMU_PER_POINT * 1.1 > available_height * 1.35


def inspect(path: Path) -> dict[str, dict[str, dict]]:
    deck = Presentation(path)
    issues: dict[str, dict[str, dict]] = {}
    for slide_index, slide in enumerate(deck.slides, 1):
        slide_issues: dict[str, dict] = {}
        shapes = list(slide.shapes)
        for index, shape in enumerate(shapes):
            if not (shape.has_text_frame or shape.shape_type == MSO_SHAPE_TYPE.PICTURE):
                continue
            bounds = box(shape)
            overflow: dict[str, str] = {}
            if bounds[0] < 0 or bounds[1] < 0 or bounds[2] > deck.slide_width or bounds[3] > deck.slide_height:
                overflow["slide"] = "shape extends beyond slide canvas"
            if text_overflow(shape):
                overflow["frame"] = "estimated text exceeds frame height"
            overlaps: list[str] = []
            for other_index, other in enumerate(shapes):
                if index == other_index or not (other.has_text_frame or other.shape_type == MSO_SHAPE_TYPE.PICTURE):
                    continue
                other_box = box(other)
                area = intersection(bounds, other_box)
                own_area = max(1, shape.width * shape.height)
                other_area = max(1, other.width * other.height)
                if area > 0 and area < min(own_area, other_area) * 0.9:
                    overlaps.append(other.name)
            if overflow or overlaps:
                entry: dict = {}
                if overflow:
                    entry["overflow"] = overflow
                if overlaps:
                    entry["overlap"] = {"overlapping_shapes": sorted(set(overlaps))}
                slide_issues[f"shape-{index + 1}"] = entry
        if slide_issues:
            issues[f"slide-{slide_index}"] = slide_issues
    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--issues-only", action="store_true")
    args = parser.parse_args()
    result = inspect(args.pptx)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
