"""Deterministic office craft findings for python-pptx decks.

The limits live here so spec revisions change data rather than scattered checks.
"""
from __future__ import annotations

import colorsys
import re

from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

OF_LIMITS = {
    "full_bleed": .92, "hue_family_degrees": 20,
    "accent_saturation": .12, "accent_light_low": .12, "accent_light_high": .93,
    "body_min_pt": 10.5, "body_max_pt": 30, "latin_measure": 90,
    "cjk_measure": 38, "narrow_measure": 10, "narrow_lines": 3,
    "numeric_share": .70, "numeric_rows": 2,
    "frame_inset_tolerance_in": .06, "frame_radius_tolerance_in": .02,
    "frame_outer_radius_min_in": .04, "group_ratio": 2,
    "empty_medium": .25, "empty_high": .35, "row_tolerance_in": .05,
}
EMU = Inches(1)
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
EMOJI = re.compile(r"[\U0001F000-\U0001F0FF\U0001F300-\U0001FAFF\u2600-\u27BF]")


def _box(shape):
    return tuple(v / EMU for v in (shape.left, shape.top, shape.width, shape.height))


def _text(shape):
    return shape.text.strip() if getattr(shape, "has_text_frame", False) else ""


def _finding(items, rule, severity, slide, detail, **extra):
    items.append({"rule": rule, "severity": severity, "slide": slide, "detail": detail, **extra})


def _numeric(value):
    raw = value.strip()
    if not raw:
        return False
    stripped = re.sub(r"[,\.\s%▲▼+\-±()$€£¥₩₹]", "", raw)
    return bool(stripped) and stripped.isdecimal()


def _radius(shape):
    try:
        return float(shape.adjustments[0]) * min(shape.width, shape.height) / EMU
    except (AttributeError, IndexError, ValueError, TypeError):
        return 0


def _gap_fraction(boxes, top, bottom):
    if bottom <= top:
        return 0
    cursor, gaps = top, []
    for _, y, _, h in sorted(boxes, key=lambda b: b[1]):
        if y + h <= top or y >= bottom:
            continue
        gaps.append(max(0, y - cursor))
        cursor = max(cursor, min(bottom, y + h))
    gaps.append(max(0, bottom - cursor))
    return max(gaps) / (bottom - top)


def analyze_craft(prs):
    """Return HIGH blockers and MEDIUM advisories; no render claims."""
    findings = []
    W, H = prs.slide_width / EMU, prs.slide_height / EMU
    for sn, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        title = next((s for s in shapes if _text(s) and _box(s)[1] < H * .22), None)
        display = len([s for s in shapes if _text(s)]) <= 2 and any(
            (r.font.size.pt if r.font.size else 0) >= OF_LIMITS["body_max_pt"]
            for s in shapes if getattr(s, "has_text_frame", False)
            for p in s.text_frame.paragraphs for r in p.runs
        )

        accents = []
        for s in shapes:
            x, y, w, h = _box(s)
            if w * h >= OF_LIMITS["full_bleed"] * W * H:
                continue
            try:
                rgb = s.fill.fore_color.rgb
                if rgb is None:
                    continue
                hue, light, sat = colorsys.rgb_to_hls(*(v / 255 for v in rgb))
                if sat >= OF_LIMITS["accent_saturation"] and OF_LIMITS["accent_light_low"] <= light <= OF_LIMITS["accent_light_high"]:
                    angle = hue * 360
                    if not any(min(abs(angle - prior), 360 - abs(angle - prior)) <= OF_LIMITS["hue_family_degrees"] for prior in accents):
                        accents.append(angle)
            except (AttributeError, TypeError):
                continue
        count = len(accents)
        if count >= (4 if display else 3):
            _finding(findings, "OF-101", "HIGH", sn, f"{count} accent families")
        elif count >= (3 if display else 2):
            _finding(findings, "OF-101", "MEDIUM", sn, f"{count} accent families")

        for s in shapes:
            if getattr(s, "has_text_frame", False):
                for p in s.text_frame.paragraphs:
                    if p._p.xpath("./a:pPr/a:buChar"):
                        char = p._p.xpath("./a:pPr/a:buChar")[0].get("char", "")
                        if EMOJI.search(char):
                            _finding(findings, "OF-108", "HIGH", sn, "emoji bullet")
                    for r in p.runs:
                        if r._r.xpath("./a:rPr/a:gradFill"):
                            _finding(findings, "OF-106", "HIGH", sn, "gradient text")
                if s is not title and _text(s):
                    sizes = [r.font.size.pt for p in s.text_frame.paragraphs for r in p.runs if r.font.size]
                    if sizes and OF_LIMITS["body_min_pt"] < max(sizes) < OF_LIMITS["body_max_pt"]:
                        line_count = sum(max(1, len(p.text.splitlines())) for p in s.text_frame.paragraphs)
                        # Explicit lines are exact; estimated wrapping is a conservative fallback.
                        width = max(.1, _box(s)[2] - (s.text_frame.margin_left + s.text_frame.margin_right) / EMU)
                        estimate = max(1, round(width * 72 / (max(sizes) * .55)))
                        line_count = max(line_count, (len(_text(s)) + estimate - 1) // estimate)
                        if line_count >= 2:
                            chars = len(_text(s).replace("\n", ""))
                            cjk = sum('\uac00' <= c <= '\ud7a3' for c in _text(s)) / max(1, chars) >= .5
                            measure = chars / line_count
                            cap = OF_LIMITS["cjk_measure" if cjk else "latin_measure"]
                            if measure > cap:
                                _finding(findings, "OF-102", "HIGH", sn, f"{measure:.1f} chars/line exceeds {cap}")
                            elif line_count >= OF_LIMITS["narrow_lines"] and measure < OF_LIMITS["narrow_measure"]:
                                _finding(findings, "OF-102", "MEDIUM", sn, f"{measure:.1f} chars/line")
            if s._element.xpath(".//a:effectLst/a:glow"):
                _finding(findings, "OF-107", "HIGH", sn, "glow effect")

            if getattr(s, "has_table", False):
                table = s.table
                # `first_row` is the table's own OOXML header flag. With no
                # header style, the first row remains the documented fallback.
                header_rows = 1 if table.first_row else 1
                rows = list(table.rows)[header_rows:]
                if len(rows) >= OF_LIMITS["numeric_rows"]:
                    for col in range(len(table.columns)):
                        values = [row.cells[col] for row in rows]
                        candidates = [c for c in values if c.text.strip()]
                        if not candidates or sum(_numeric(c.text) for c in candidates) / len(candidates) < OF_LIMITS["numeric_share"]:
                            continue
                        for c in values:
                            if not _numeric(c.text):
                                continue
                            for p in c.text_frame.paragraphs:
                                if p.alignment is None:
                                    _finding(findings, "OF-103", "MEDIUM", sn, "numeric alignment inherited", tier="derived")
                                elif p.alignment != PP_ALIGN.RIGHT:
                                    _finding(findings, "OF-103", "HIGH", sn, "numeric cell not right aligned")

        rounds = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and s.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE]
        for outer in rounds:
            ox, oy, ow, oh = _box(outer)
            for inner in rounds:
                if inner is outer:
                    continue
                ix, iy, iw, ih = _box(inner)
                pads = (ix-ox, iy-oy, ox+ow-ix-iw, oy+oh-iy-ih)
                if min(pads) < 0 or max(pads)-min(pads) > OF_LIMITS["frame_inset_tolerance_in"]:
                    continue
                radius = _radius(outer)
                expected = max(0, radius - sum(pads)/4)
                if radius > OF_LIMITS["frame_outer_radius_min_in"] and abs(_radius(inner)-expected) > OF_LIMITS["frame_radius_tolerance_in"]:
                    _finding(findings, "OF-104", "HIGH", sn, "nested frame radius differs from inset")

        # Explicit card boundaries make group membership measurable. Until the
        # template corpus is recalibrated, OF-105 remains advisory.
        cards = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and
                 _box(s)[2] > 1 and _box(s)[3] > .7 and s is not title]
        groups = []
        for card in cards:
            cx, cy, cw, ch = _box(card)
            members = [s for s in shapes if s is not card and _text(s) and
                       cx <= _box(s)[0] and cy <= _box(s)[1] and
                       _box(s)[0]+_box(s)[2] <= cx+cw and _box(s)[1]+_box(s)[3] <= cy+ch]
            if len(members) >= 2:
                member_boxes = [_box(s) for s in members]
                gaps = []
                for index, a in enumerate(member_boxes):
                    for b in member_boxes[index+1:]:
                        dx = max(0, b[0]-a[0]-a[2], a[0]-b[0]-b[2])
                        dy = max(0, b[1]-a[1]-a[3], a[1]-b[1]-b[3])
                        gaps.append((dx*dx+dy*dy)**.5)
                groups.append((card, min(gaps, default=0), members))
        for i, (a, gap_a, _) in enumerate(groups):
            ax, ay, aw, ah = _box(a)
            for b, gap_b, _ in groups[i+1:]:
                bx, by, bw, bh = _box(b)
                if abs(ay-by) > .1 or abs(ah-bh) > .1:
                    continue
                between = max(0, bx-ax-aw) if ax < bx else max(0, ax-bx-bw)
                if between < OF_LIMITS["group_ratio"] * max(gap_a, gap_b):
                    _finding(findings, "OF-105", "MEDIUM", sn, "between-card gap below 2x within-card gap")

        if not display:
            top = (_box(title)[1] + _box(title)[3] + .1) if title else .5
            bottom = H - .55
            content = [_box(s) for s in shapes if s is not title and
                       (_text(s) or s in cards or getattr(s, "has_table", False) or s.shape_type == MSO_SHAPE_TYPE.PICTURE) and
                       _box(s)[1] >= top and _box(s)[1] < bottom]
            fraction = _gap_fraction(content, top, bottom)
            if fraction > OF_LIMITS["empty_high"]:
                _finding(findings, "OF-109", "HIGH", sn, f"largest empty band {fraction:.2f}")
            elif fraction >= OF_LIMITS["empty_medium"]:
                _finding(findings, "OF-109", "MEDIUM", sn, f"largest empty band {fraction:.2f}")

            # Equal-height card rows inherit their height from their fullest
            # sibling. Score the row's least-empty card, once per row.
            card_fractions = []
            for card in cards:
                cx, cy, cw, ch = _box(card)
                if cw * ch >= OF_LIMITS["full_bleed"] * W * H:
                    continue
                children = [s for s in shapes if s is not card and s is not title and
                            cx <= _box(s)[0] and cy <= _box(s)[1] and
                            _box(s)[0]+_box(s)[2] <= cx+cw and _box(s)[1]+_box(s)[3] <= cy+ch and _text(s)]
                if not children:
                    continue
                biggest = max((r.font.size.pt if r.font.size else 0 for s in children if getattr(s, "has_text_frame", False)
                               for p in s.text_frame.paragraphs for r in p.runs), default=0)
                if biggest >= OF_LIMITS["body_max_pt"] and len(children) <= 2:
                    continue
                margin_top = card.text_frame.margin_top / EMU if card.has_text_frame else 0
                margin_bottom = card.text_frame.margin_bottom / EMU if card.has_text_frame else 0
                inner_top, inner_bottom = cy + margin_top, cy + ch - margin_bottom
                card_fractions.append((card, _gap_fraction([_box(child) for child in children], inner_top, inner_bottom)))
            seen_cards = set()
            for card, value in card_fractions:
                if id(card) in seen_cards:
                    continue
                cx, cy, cw, ch = _box(card)
                row = [(other, amount) for other, amount in card_fractions if
                       abs(_box(other)[1]-cy) <= OF_LIMITS["row_tolerance_in"] and
                       abs(_box(other)[3]-ch) <= OF_LIMITS["row_tolerance_in"]]
                seen_cards.update(id(other) for other, _ in row)
                fraction = min(amount for _, amount in row)
                if fraction > OF_LIMITS["empty_high"]:
                    _finding(findings, "OF-109", "HIGH", sn, f"card-row empty band {fraction:.2f}")
                elif fraction >= OF_LIMITS["empty_medium"]:
                    _finding(findings, "OF-109", "MEDIUM", sn, f"card-row empty band {fraction:.2f}")

    return {"pass": not any(f["severity"] == "HIGH" for f in findings), "findings": findings}
