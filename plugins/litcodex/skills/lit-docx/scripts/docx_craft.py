"""The lit-docx gate: office craft checks over the DOCX XML, plus the document output checks.

`audit_docx_office_floor` (OF-301 measure, OF-302 numeric alignment) reads the XML with the
standard library. Run as a script, this file is the document gate:

    python3 docx_craft.py out.docx --source draft.md [--layout] [--tonality Report] [--compare other.docx]

It adds a content check (every Markdown heading reached the document, no unfilled blank, no
frontmatter printed as text) and docx_layout.py: noun-phrase headings and heading order always;
the restraint checks of the document spec on a tonality document (date.iso on every path);
with --layout the pages, rendered through soffice in a private profile, for fill, stranded
headings, splits, component variety and the cover block; with --compare the structural
difference from the same source in another tonality. Prints JSON; exit 0 on PASS, 1 on FAIL.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
OF_LIMITS = {
    'latin_measure': 90, 'cjk_measure': 38, 'cjk_majority': .5,
    'latin_glyph_width': .55, 'cjk_glyph_width': 1.0,
    'default_font_pt': 11, 'numeric_share': .70, 'numeric_min_rows': 2,
}


@dataclass
class DocxFinding:
    rule_id: str
    line: int
    column: int
    section: str
    message: str
    excerpt: str
    severity: str
    tier: str = 'derived'


def _text(node):
    return ''.join(part.text or '' for part in node.iter(f'{W}t')).strip()


def _numeric(text):
    stripped = re.sub(r'[,\.\s%▲▼+\-±()$€£¥₩₹]', '', text.strip())
    return bool(stripped) and stripped.isdecimal()


def audit_docx_office_floor(path: Path) -> list[DocxFinding]:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
        try:
            styles = ET.fromstring(archive.read('word/styles.xml'))
        except KeyError:
            styles = None
    body = root.find(f'{W}body')
    if body is None:
        return []
    result = []
    section = body.find(f'{W}sectPr')
    page = section.find(f'{W}pgSz') if section is not None else None
    margins = section.find(f'{W}pgMar') if section is not None else None
    if page is not None and margins is not None:
        width = (int(page.get(f'{W}w', '12240')) - int(margins.get(f'{W}left', '1440')) - int(margins.get(f'{W}right', '1440'))) / 1440
    else:
        width = 6.5
    normal_size = OF_LIMITS['default_font_pt']
    if styles is not None:
        for style in styles.findall(f'{W}style'):
            if style.get(f'{W}styleId') == 'Normal':
                size = style.find(f'.//{W}sz')
                if size is not None:
                    normal_size = int(size.get(f'{W}val', '22')) / 2
                break
    for paragraph in body.findall(f'{W}p'):
        text = _text(paragraph)
        if not text:
            continue
        sizes = [int(node.get(f'{W}val')) / 2 for node in paragraph.findall(f'.//{W}sz') if node.get(f'{W}val')]
        font_pt = max(sizes, default=normal_size)
        cjk = sum('\uac00' <= char <= '\ud7a3' for char in text) / len(text) >= OF_LIMITS['cjk_majority']
        estimate = width * 72 / (font_pt * OF_LIMITS['cjk_glyph_width' if cjk else 'latin_glyph_width'])
        ceiling = OF_LIMITS['cjk_measure' if cjk else 'latin_measure']
        if estimate > ceiling:
            result.append(DocxFinding('OF-301', 0, 0, 'DOCX', f'Estimated {estimate:.1f} characters per line exceeds {ceiling}', text[:80], 'MEDIUM'))
    for table in body.findall(f'{W}tbl'):
        rows = table.findall(f'{W}tr')[1:]
        if len(rows) < OF_LIMITS['numeric_min_rows']:
            continue
        columns = max((len(row.findall(f'{W}tc')) for row in rows), default=0)
        for index in range(columns):
            cells = [row.findall(f'{W}tc')[index] for row in rows if len(row.findall(f'{W}tc')) > index]
            candidates = [cell for cell in cells if _text(cell)]
            if not candidates or sum(_numeric(_text(cell)) for cell in candidates) / len(candidates) < OF_LIMITS['numeric_share']:
                continue
            for cell in cells:
                if not _numeric(_text(cell)):
                    continue
                for paragraph in cell.findall(f'{W}p'):
                    align = paragraph.find(f'./{W}pPr/{W}jc')
                    if align is None:
                        result.append(DocxFinding('OF-302', 0, index + 1, 'DOCX', 'Numeric alignment inherited; not verified', _text(cell)[:80], 'MEDIUM'))
                    elif align.get(f'{W}val') != 'right':
                        result.append(DocxFinding('OF-302', 0, index + 1, 'DOCX', 'Numeric cell is explicitly not right aligned', _text(cell)[:80], 'HIGH', 'measured'))
    return result


HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
BLANK = re.compile(r"\[(?![^\]]*\d)[^\]\n]{1,12}\]|\bXXX\b|\bTBD\b|○○|\{\{[^}]*\}\}")
FRONTMATTER_LINE = re.compile(r"^(title|subtitle|author|organization|date|notice|tonality|publisher|lang):\s")


def check_content(docx_path: Path, source_text: str | None) -> dict:
    """Body text exists, every source heading reached the page, no blank or frontmatter left."""
    import docx_layout
    from docx import Document

    document = Document(str(docx_path))
    paragraphs = [p.text for p in document.paragraphs]
    cells = [c.text for t in document.tables for r in t.rows for c in r.cells]
    page_text = docx_layout.norm("".join(paragraphs + cells))
    missing, fenced = [], False
    for line in (source_text or "").split("---", 2)[-1].splitlines() if source_text else []:
        if line.lstrip().startswith(("```", "~~~")):
            fenced = not fenced
            continue
        match = None if fenced else HEADING.match(line)
        if match and docx_layout.norm(re.sub(r"[*_`]", "", match.group(2))) not in page_text:
            missing.append(re.sub(r"[*_`]", "", match.group(2)).strip())
    blanks = sorted({m.group(0) for text in paragraphs + cells for m in BLANK.finditer(text)})
    leaked = [p for p in paragraphs[:8] if FRONTMATTER_LINE.match(p)]
    chars = sum(len(p.strip()) for p in paragraphs + cells)
    return {"pass": chars > 0 and not missing and not blanks and not leaked, "body_chars": chars,
            "missing_headings": missing, "blanks": blanks, "frontmatter_leak": leaked}


def render_pdf(docx_path: Path, out_dir: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory(prefix="lit-docx-profile-") as profile:
        subprocess.run([soffice, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", "pdf",
                        "--outdir", str(out_dir), str(docx_path)], capture_output=True, text=True, timeout=300)
    pdf = out_dir / f"{docx_path.stem}.pdf"
    return pdf if pdf.exists() else None


def output_checks(docx_path: Path, front: dict, args) -> dict:
    import docx_layout

    tonality = bool(args.tonality or front.get("tonality"))
    info = docx_layout.read_docx(docx_path)
    found = docx_layout.structural(info, front, args.kind == "manuscript", tonality) + docx_layout.restraint(docx_path, tonality)
    result = {"components": info["components"], "rendered": False}
    if args.compare:
        mine, other = docx_layout.structure(docx_path), docx_layout.structure(args.compare.expanduser().resolve())
        differ = docx_layout.structure_diff(mine, other)
        result["structure"] = {"this": mine, "other": other, "differ": differ}
        if len(differ) < 3:
            found.append(docx_layout.finding(
                "tonality.structure", "FAIL", None, f"{len(differ)} structural difference(s): {', '.join(differ) or 'none'} (at least 3)",
                "Tonalities differ in structure: title block, summary form, numbering, components, running head, contents."))
    if args.layout:
        with tempfile.TemporaryDirectory(prefix="lit-docx-pages-") as tmp:
            pdf = render_pdf(docx_path, Path(tmp))
            if pdf is None:
                result["render"] = "soffice missing or the conversion failed: the page checks did not run"
            else:
                page_found, pages = docx_layout.paged(info, docx_layout.read_pdf(pdf), tonality, front)
                found += page_found
                result.update(pages, rendered=True)
    result["findings"] = [f for f in found if f["severity"] == "FAIL"]
    result["advisories"] = [f for f in found if f["severity"] != "FAIL"]
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Pass/fail gate for a generated DOCX.")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--source", type=Path, help="the Markdown the DOCX was converted from")
    parser.add_argument("--kind", choices=("report", "manuscript"), default="report")
    parser.add_argument("--layout", action="store_true", help="render through soffice and check the pages")
    parser.add_argument("--tonality", help="the tonality used, when the source frontmatter does not name it")
    parser.add_argument("--compare", type=Path, help="the same source in another tonality (structure check)")
    args = parser.parse_args(argv)
    sys.dont_write_bytecode = True
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from convert_md_to_docx import parse_frontmatter

    docx_path = args.docx.expanduser().resolve()
    source_text = args.source.read_text(encoding="utf-8") if args.source else None
    front = parse_frontmatter(source_text)[0] if source_text and source_text.startswith("---") else {}
    floor = audit_docx_office_floor(docx_path)
    content = check_content(docx_path, source_text)
    output = output_checks(docx_path, front, args)
    blocking = [f for f in floor if f.severity == "HIGH"]
    report = {
        "file": str(docx_path),
        "pass": not blocking and content["pass"] and not output["findings"],
        "office_floor": [f.__dict__ for f in floor],
        "content": content,
        "output": output,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
